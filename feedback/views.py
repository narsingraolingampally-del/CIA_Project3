from django.contrib import messages
from django.contrib.auth.decorators import login_required, user_passes_test
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

import openpyxl

from quiz.models import StudentProfile

from .forms import (
    FeedbackFormForm,
    FeedbackQuestionForm,
    FeedbackQuestionUploadForm,
    FeedbackAssignmentForm,
)

from .models import (
    FeedbackForm,
    FeedbackAssignment,
    FeedbackQuestion,
    FeedbackResponse,
    FeedbackAnswer,
)


# ============================================================
# ADMIN ACCESS
# ============================================================

def is_admin(user):
    return user.is_authenticated and user.is_superuser


def admin_required(view_func):
    return login_required(
        login_url="/admin-login/"
    )(
        user_passes_test(
            is_admin,
            login_url="/admin-login/"
        )(view_func)
    )


# ============================================================
# ADMIN FEEDBACK DASHBOARD
# ============================================================

@admin_required
def feedback_admin_dashboard(request):

    total_forms = FeedbackForm.objects.count()

    active_forms = FeedbackForm.objects.filter(
        is_active=True
    ).count()

    total_assignments = FeedbackAssignment.objects.count()

    total_responses = FeedbackResponse.objects.count()

    context = {
        "total_forms": total_forms,
        "active_forms": active_forms,
        "total_assignments": total_assignments,
        "total_responses": total_responses,
    }

    return render(
        request,
        "feedback/admin_dashboard.html",
        context,
    )


# ============================================================
# FEEDBACK FORM MANAGEMENT
# ============================================================

@admin_required
def feedback_forms(request):

    forms = FeedbackForm.objects.select_related(
        "created_by"
    ).order_by("-created_at")

    return render(
        request,
        "feedback/forms.html",
        {
            "feedback_forms": forms,
        },
    )


@admin_required
def create_feedback_form(request):

    if request.method == "POST":

        form = FeedbackFormForm(request.POST)

        if form.is_valid():

            feedback_form = form.save(commit=False)

            feedback_form.created_by = request.user

            # New forms are inactive until Admin activates them.
            feedback_form.is_active = False

            feedback_form.save()

            messages.success(
                request,
                f"Feedback form '{feedback_form.title}' created successfully."
            )

            return redirect("feedback_forms")

    else:

        form = FeedbackFormForm()

    return render(
        request,
        "feedback/create_form.html",
        {
            "form": form,
        },
    )


@admin_required
def edit_feedback_form(request, pk):

    feedback_form = get_object_or_404(
        FeedbackForm,
        pk=pk,
    )

    if request.method == "POST":

        form = FeedbackFormForm(
            request.POST,
            instance=feedback_form,
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                f"Feedback form '{feedback_form.title}' updated successfully."
            )

            return redirect("feedback_forms")

    else:

        form = FeedbackFormForm(
            instance=feedback_form,
        )

    return render(
        request,
        "feedback/edit_form.html",
        {
            "form": form,
            "feedback_form": feedback_form,
        },
    )


@admin_required
def toggle_feedback_form(request, pk):

    feedback_form = get_object_or_404(
        FeedbackForm,
        pk=pk,
    )

    if request.method != "POST":
        return redirect("feedback_forms")

    feedback_form.is_active = not feedback_form.is_active

    feedback_form.save(
        update_fields=[
            "is_active",
            "updated_at",
        ]
    )

    if feedback_form.is_active:

        messages.success(
            request,
            f"'{feedback_form.title}' has been activated."
        )

    else:

        messages.warning(
            request,
            f"'{feedback_form.title}' has been deactivated."
        )

    return redirect("feedback_forms")


@admin_required
def delete_feedback_form(request, pk):

    feedback_form = get_object_or_404(
        FeedbackForm,
        pk=pk,
    )

    if request.method != "POST":
        return redirect("feedback_forms")

    title = feedback_form.title

    feedback_form.delete()

    messages.success(
        request,
        f"Feedback form '{title}' deleted successfully."
    )

    return redirect("feedback_forms")


# ============================================================
# FEEDBACK QUESTIONS
# ============================================================

@admin_required
def feedback_questions(request, form_id):

    feedback_form = get_object_or_404(
        FeedbackForm,
        pk=form_id,
    )

    questions = feedback_form.questions.all().order_by(
        "display_order",
        "id",
    )

    return render(
        request,
        "feedback/questions.html",
        {
            "feedback_form": feedback_form,
            "questions": questions,
        },
    )


@admin_required
def add_feedback_question(request, form_id):

    feedback_form = get_object_or_404(
        FeedbackForm,
        pk=form_id,
    )

    if request.method == "POST":

        form = FeedbackQuestionForm(request.POST)

        if form.is_valid():

            question = form.save(
                commit=False
            )

            question.feedback_form = feedback_form

            question.save()

            messages.success(
                request,
                "Feedback question added successfully.",
            )

            return redirect(
                "feedback_questions",
                form_id=feedback_form.id,
            )

    else:

        next_order = (
            feedback_form.questions.count() + 1
        )

        form = FeedbackQuestionForm(
            initial={
                "display_order": next_order,
            }
        )

    return render(
        request,
        "feedback/add_question.html",
        {
            "form": form,
            "feedback_form": feedback_form,
        },
    )


@admin_required
def edit_feedback_question(request, pk):

    question = get_object_or_404(
        FeedbackQuestion,
        pk=pk,
    )

    if request.method == "POST":

        form = FeedbackQuestionForm(
            request.POST,
            instance=question,
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Feedback question updated successfully.",
            )

            return redirect(
                "feedback_questions",
                form_id=question.feedback_form.id,
            )

    else:

        form = FeedbackQuestionForm(
            instance=question,
        )

    return render(
        request,
        "feedback/edit_question.html",
        {
            "form": form,
            "question": question,
            "feedback_form": question.feedback_form,
        },
    )


@admin_required
def toggle_feedback_question(request, pk):

    question = get_object_or_404(
        FeedbackQuestion,
        pk=pk,
    )

    if request.method != "POST":

        return redirect(
            "feedback_questions",
            form_id=question.feedback_form.id,
        )

    question.is_active = not question.is_active

    question.save(
        update_fields=[
            "is_active"
        ]
    )

    if question.is_active:

        messages.success(
            request,
            "Feedback question activated.",
        )

    else:

        messages.warning(
            request,
            "Feedback question deactivated.",
        )

    return redirect(
        "feedback_questions",
        form_id=question.feedback_form.id,
    )


@admin_required
def delete_feedback_question(request, pk):

    question = get_object_or_404(
        FeedbackQuestion,
        pk=pk,
    )

    form_id = question.feedback_form.id

    if request.method != "POST":

        return redirect(
            "feedback_questions",
            form_id=form_id,
        )

    question.delete()

    messages.success(
        request,
        "Feedback question deleted successfully.",
    )

    return redirect(
        "feedback_questions",
        form_id=form_id,
    )


# ============================================================
# DOWNLOAD FEEDBACK QUESTION EXCEL TEMPLATE
# ============================================================

@admin_required
def download_feedback_question_template(request, form_id):

    feedback_form = get_object_or_404(
        FeedbackForm,
        pk=form_id,
    )

    workbook = openpyxl.Workbook()

    worksheet = workbook.active
    worksheet.title = "Questions"

    headers = [
        "question_text",
        "question_type",
        "option1",
        "option2",
        "option3",
        "option4",
        "display_order",
        "is_required",
    ]

    worksheet.append(headers)

    worksheet.append([
        "Faculty explains concepts clearly.",
        "rating",
        "",
        "",
        "",
        "",
        1,
        "yes",
    ])

    worksheet.append([
        "Faculty encourages student participation.",
        "rating",
        "",
        "",
        "",
        "",
        2,
        "yes",
    ])

    worksheet.append([
        "The faculty uses appropriate teaching methods.",
        "rating",
        "",
        "",
        "",
        "",
        3,
        "yes",
    ])

    worksheet.append([
        "How satisfied are you overall?",
        "choice",
        "Very Satisfied",
        "Satisfied",
        "Neutral",
        "Not Satisfied",
        4,
        "yes",
    ])

    worksheet.append([
        "Additional Comments",
        "text",
        "",
        "",
        "",
        "",
        5,
        "no",
    ])

    worksheet.column_dimensions["A"].width = 55
    worksheet.column_dimensions["B"].width = 18
    worksheet.column_dimensions["C"].width = 25
    worksheet.column_dimensions["D"].width = 25
    worksheet.column_dimensions["E"].width = 25
    worksheet.column_dimensions["F"].width = 25
    worksheet.column_dimensions["G"].width = 15
    worksheet.column_dimensions["H"].width = 15

    instructions = workbook.create_sheet(
        "Instructions"
    )

    instructions.append([
        "Feedback Question Upload Instructions"
    ])

    instructions.append([])

    instructions.append([
        "question_type values:",
        "rating, yes_no, text, choice",
    ])

    instructions.append([
        "rating:",
        "Student selects 1 to 5",
    ])

    instructions.append([
        "yes_no:",
        "Student selects Yes or No",
    ])

    instructions.append([
        "text:",
        "Student enters written response",
    ])

    instructions.append([
        "choice:",
        "Use option1 to option4",
    ])

    instructions.append([
        "is_required:",
        "yes/no, true/false, 1/0",
    ])

    instructions.append([])

    instructions.append([
        "Feedback Form:",
        feedback_form.title,
    ])

    instructions.column_dimensions["A"].width = 30
    instructions.column_dimensions["B"].width = 70

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; filename="feedback_questions_template.xlsx"'
    )

    workbook.save(response)

    return response


# ============================================================
# BULK UPLOAD FEEDBACK QUESTIONS
# ============================================================

@admin_required
def upload_feedback_questions(request, form_id):

    feedback_form = get_object_or_404(
        FeedbackForm,
        pk=form_id,
    )

    if request.method != "POST":

        form = FeedbackQuestionUploadForm()

        return render(
            request,
            "feedback/upload_questions.html",
            {
                "form": form,
                "feedback_form": feedback_form,
            },
        )

    form = FeedbackQuestionUploadForm(
        request.POST,
        request.FILES,
    )

    if not form.is_valid():

        return render(
            request,
            "feedback/upload_questions.html",
            {
                "form": form,
                "feedback_form": feedback_form,
            },
        )

    uploaded_file = form.cleaned_data[
        "excel_file"
    ]

    try:

        workbook = openpyxl.load_workbook(
            uploaded_file,
            read_only=True,
            data_only=True,
        )

        worksheet = workbook.active

        rows = list(
            worksheet.iter_rows(
                values_only=True
            )
        )

        if not rows:

            raise ValueError(
                "The Excel file is empty."
            )

        headers = [
            str(value).strip().lower()
            if value is not None
            else ""
            for value in rows[0]
        ]

        required_headers = [
            "question_text",
            "question_type",
            "option1",
            "option2",
            "option3",
            "option4",
            "display_order",
            "is_required",
        ]

        missing_headers = [
            header
            for header in required_headers
            if header not in headers
        ]

        if missing_headers:

            raise ValueError(
                "Missing required columns: "
                + ", ".join(missing_headers)
            )

        column_index = {
            header: headers.index(header)
            for header in required_headers
        }

        valid_types = {
            "rating",
            "yes_no",
            "text",
            "choice",
        }

        validated_questions = []
        errors = []

        base_order = feedback_form.questions.count()

        for excel_row_number, row in enumerate(
            rows[1:],
            start=2,
        ):

            def get_value(column_name):

                index = column_index[
                    column_name
                ]

                if index >= len(row):
                    return None

                return row[index]

            question_text = get_value(
                "question_text"
            )

            question_type = get_value(
                "question_type"
            )

            option1 = get_value("option1")
            option2 = get_value("option2")
            option3 = get_value("option3")
            option4 = get_value("option4")

            display_order = get_value(
                "display_order"
            )

            is_required = get_value(
                "is_required"
            )

            question_text = (
                str(question_text).strip()
                if question_text is not None
                else ""
            )

            question_type = (
                str(question_type).strip().lower()
                if question_type is not None
                else ""
            )

            if not question_text:

                errors.append(
                    f"Row {excel_row_number}: "
                    "question_text is required."
                )

                continue

            if question_type not in valid_types:

                errors.append(
                    f"Row {excel_row_number}: invalid "
                    f"question_type '{question_type}'."
                )

                continue

            if display_order in (
                None,
                "",
            ):

                display_order = (
                    base_order
                    + len(validated_questions)
                    + 1
                )

            else:

                try:

                    display_order = int(
                        float(display_order)
                    )

                    if display_order < 1:
                        raise ValueError

                except (
                    TypeError,
                    ValueError,
                ):

                    errors.append(
                        f"Row {excel_row_number}: "
                        "display_order must be a positive number."
                    )

                    continue

            if isinstance(
                is_required,
                bool,
            ):

                required_value = is_required

            else:

                required_text = (
                    str(is_required).strip().lower()
                    if is_required is not None
                    else "yes"
                )

                if required_text in {
                    "yes",
                    "true",
                    "1",
                    "required",
                }:

                    required_value = True

                elif required_text in {
                    "no",
                    "false",
                    "0",
                    "optional",
                }:

                    required_value = False

                else:

                    errors.append(
                        f"Row {excel_row_number}: "
                        "is_required must be yes/no, "
                        "true/false, or 1/0."
                    )

                    continue

            options = [
                option1,
                option2,
                option3,
                option4,
            ]

            options = [
                str(option).strip()
                if option is not None
                else ""
                for option in options
            ]

            if question_type == "choice":

                if not any(options):

                    errors.append(
                        f"Row {excel_row_number}: multiple-choice "
                        "questions require at least one option."
                    )

                    continue

            else:

                options = [
                    None,
                    None,
                    None,
                    None,
                ]

            validated_questions.append({
                "question_text": question_text,
                "question_type": question_type,
                "option1": options[0],
                "option2": options[1],
                "option3": options[2],
                "option4": options[3],
                "display_order": display_order,
                "is_required": required_value,
            })

        workbook.close()

        if errors:

            return render(
                request,
                "feedback/upload_questions.html",
                {
                    "form": form,
                    "feedback_form": feedback_form,
                    "errors": errors,
                },
            )

        if not validated_questions:

            return render(
                request,
                "feedback/upload_questions.html",
                {
                    "form": form,
                    "feedback_form": feedback_form,
                    "errors": [
                        "The Excel file does not contain any question rows."
                    ],
                },
            )

        with transaction.atomic():

            FeedbackQuestion.objects.bulk_create([
                FeedbackQuestion(
                    feedback_form=feedback_form,
                    question_text=item[
                        "question_text"
                    ],
                    question_type=item[
                        "question_type"
                    ],
                    option1=item["option1"],
                    option2=item["option2"],
                    option3=item["option3"],
                    option4=item["option4"],
                    display_order=item[
                        "display_order"
                    ],
                    is_required=item[
                        "is_required"
                    ],
                    is_active=True,
                )
                for item in validated_questions
            ])

        messages.success(
            request,
            f"{len(validated_questions)} feedback "
            "questions uploaded successfully.",
        )

        return redirect(
            "feedback_questions",
            form_id=feedback_form.id,
        )

    except Exception as exc:

        return render(
            request,
            "feedback/upload_questions.html",
            {
                "form": form,
                "feedback_form": feedback_form,
                "errors": [
                    f"Unable to process the Excel file: {exc}"
                ],
            },
        )


# ============================================================
# FEEDBACK ASSIGNMENTS
# ============================================================

@admin_required
def feedback_assignments(request):

    assignments = FeedbackAssignment.objects.select_related(
        "feedback_form",
        "course",
        "subject",
        "faculty",
        "faculty__user",
    ).order_by(
        "-created_at"
    )

    return render(
        request,
        "feedback/assignments.html",
        {
            "assignments": assignments,
        },
    )


@admin_required
def create_feedback_assignment(request):

    feedback_forms = FeedbackForm.objects.order_by(
        "title"
    )

    if request.method == "POST":

        form = FeedbackAssignmentForm(
            request.POST
        )

        feedback_form_id = request.POST.get(
            "feedback_form"
        )

        if not feedback_form_id:

            messages.error(
                request,
                "Please select a Feedback Form."
            )

        elif form.is_valid():

            feedback_form = get_object_or_404(
                FeedbackForm,
                pk=feedback_form_id,
            )

            assignment = form.save(
                commit=False
            )

            assignment.feedback_form = (
                feedback_form
            )

            assignment.save()

            messages.success(
                request,
                "Feedback assignment created successfully.",
            )

            return redirect(
                "feedback_assignments"
            )

        else:

            print(
                "FEEDBACK ASSIGNMENT FORM ERRORS:"
            )

            print(form.errors)

    else:

        form = FeedbackAssignmentForm()

    return render(
        request,
        "feedback/create_assignment.html",
        {
            "form": form,
            "feedback_forms": feedback_forms,
        },
    )


@admin_required
def toggle_feedback_assignment(request, pk):

    assignment = get_object_or_404(
        FeedbackAssignment,
        pk=pk,
    )

    if request.method != "POST":

        return redirect(
            "feedback_assignments"
        )

    assignment.is_active = not assignment.is_active

    assignment.save(
        update_fields=[
            "is_active"
        ]
    )

    if assignment.is_active:

        messages.success(
            request,
            "Feedback assignment activated.",
        )

    else:

        messages.warning(
            request,
            "Feedback assignment deactivated.",
        )

    return redirect(
        "feedback_assignments"
    )


@admin_required
def delete_feedback_assignment(request, pk):

    assignment = get_object_or_404(
        FeedbackAssignment,
        pk=pk,
    )

    if request.method != "POST":

        return redirect(
            "feedback_assignments"
        )

    assignment.delete()

    messages.success(
        request,
        "Feedback assignment deleted successfully.",
    )

    return redirect(
        "feedback_assignments"
    )


# ============================================================
# STUDENT ACCESS
# ============================================================

def is_student(user):

    return (
        user.is_authenticated
        and user.is_student
    )


def student_required(view_func):

    return login_required(
        login_url="/student/login/"
    )(
        user_passes_test(
            is_student,
            login_url="/student/login/"
        )(view_func)
    )


# ============================================================
# STUDENT FEEDBACK LIST
# ============================================================

@student_required
def student_feedback_list(request):

    student = get_object_or_404(
        StudentProfile,
        user=request.user,
    )

    now = timezone.now()

    assignments = (
        FeedbackAssignment.objects
        .select_related(
            "feedback_form",
            "course",
            "subject",
            "faculty",
            "faculty__user",
        )
        .filter(
            is_active=True,
            feedback_form__is_active=True,
            course=student.course,
            semester=student.semester,
            academic_year=student.academic_year,
        )
        .order_by(
            "start_date",
            "id",
        )
    )

    available_assignments = []

    for assignment in assignments:

        if (
            assignment.start_date
            and now < assignment.start_date
        ):

            continue

        if (
            assignment.end_date
            and now > assignment.end_date
        ):

            continue

        already_submitted = (
            FeedbackResponse.objects.filter(
                assignment=assignment,
                student=student,
            ).exists()
        )

        available_assignments.append({
            "assignment": assignment,
            "already_submitted": already_submitted,
        })

    return render(
        request,
        "feedback/student_feedback_list.html",
        {
            "student": student,
            "available_assignments": (
                available_assignments
            ),
        },
    )


# ============================================================
# STUDENT FEEDBACK FORM
# ============================================================

@student_required
def student_feedback_form(
    request,
    assignment_id,
):

    student = get_object_or_404(
        StudentProfile,
        user=request.user,
    )

    assignment = get_object_or_404(
        FeedbackAssignment.objects.select_related(
            "feedback_form",
            "course",
            "subject",
            "faculty",
            "faculty__user",
        ),
        pk=assignment_id,
        is_active=True,
        feedback_form__is_active=True,
    )

    # Course check
    if assignment.course_id != student.course_id:

        messages.error(
            request,
            "This feedback is not available for your course.",
        )

        return redirect(
            "student_feedback_list"
        )

    # Semester check
    if assignment.semester != student.semester:

        messages.error(
            request,
            "This feedback is not available for your semester.",
        )

        return redirect(
            "student_feedback_list"
        )

    # Academic year check
    if (
        assignment.academic_year
        != student.academic_year
    ):

        messages.error(
            request,
            "This feedback is not available for your academic year.",
        )

        return redirect(
            "student_feedback_list"
        )

    now = timezone.now()

    # Start date
    if (
        assignment.start_date
        and now < assignment.start_date
    ):

        messages.warning(
            request,
            "This feedback has not started yet.",
        )

        return redirect(
            "student_feedback_list"
        )

    # End date
    if (
        assignment.end_date
        and now > assignment.end_date
    ):

        messages.warning(
            request,
            "This feedback is closed.",
        )

        return redirect(
            "student_feedback_list"
        )

    # Prevent duplicate submission
    existing_response = (
        FeedbackResponse.objects.filter(
            assignment=assignment,
            student=student,
        ).first()
    )

    if existing_response:

        return render(
            request,
            "feedback/student_feedback_success.html",
            {
                "assignment": assignment,
                "response": existing_response,
                "already_submitted": True,
            },
        )

    # Get active questions
    questions = (
        FeedbackQuestion.objects.filter(
            feedback_form=assignment.feedback_form,
            is_active=True,
        )
        .order_by(
            "display_order",
            "id",
        )
    )

    if not questions.exists():

        messages.error(
            request,
            "This feedback form does not contain any active questions.",
        )

        return redirect(
            "student_feedback_list"
        )

    # --------------------------------------------------------
    # SUBMISSION
    # --------------------------------------------------------

    if request.method == "POST":

        validation_errors = []
        submitted_answers = []

        for question in questions:

            field_name = (
                f"question_{question.id}"
            )

            value = request.POST.get(
                field_name,
                "",
            )

            value = value.strip()

            # Required question
            if (
                question.is_required
                and not value
            ):

                validation_errors.append(
                    f"Question {question.display_order} is required."
                )

                continue

            # Optional blank answer
            if not value:

                submitted_answers.append({
                    "question": question,
                    "rating": None,
                    "answer_text": None,
                    "selected_option": None,
                })

                continue

            # ------------------------------------------------
            # RATING
            # ------------------------------------------------

            if question.question_type == "rating":

                try:

                    rating = int(value)

                except (
                    TypeError,
                    ValueError,
                ):

                    validation_errors.append(
                        f"Question {question.display_order}: "
                        "rating must be between 1 and 5."
                    )

                    continue

                if rating < 1 or rating > 5:

                    validation_errors.append(
                        f"Question {question.display_order}: "
                        "rating must be between 1 and 5."
                    )

                    continue

                submitted_answers.append({
                    "question": question,
                    "rating": rating,
                    "answer_text": None,
                    "selected_option": None,
                })

            # ------------------------------------------------
            # YES / NO
            # ------------------------------------------------

            elif question.question_type == "yes_no":

                normalized = value.lower()

                if normalized not in {
                    "yes",
                    "no",
                }:

                    validation_errors.append(
                        f"Question {question.display_order}: "
                        "please select Yes or No."
                    )

                    continue

                submitted_answers.append({
                    "question": question,
                    "rating": None,
                    "answer_text": None,
                    "selected_option": normalized,
                })

            # ------------------------------------------------
            # CHOICE
            # ------------------------------------------------

            elif question.question_type == "choice":

                valid_options = [
                    option.strip()
                    for option in [
                        question.option1,
                        question.option2,
                        question.option3,
                        question.option4,
                    ]
                    if option
                ]

                if value not in valid_options:

                    validation_errors.append(
                        f"Question {question.display_order}: "
                        "invalid option selected."
                    )

                    continue

                submitted_answers.append({
                    "question": question,
                    "rating": None,
                    "answer_text": None,
                    "selected_option": value,
                })

            # ------------------------------------------------
            # TEXT
            # ------------------------------------------------

            elif question.question_type == "text":

                submitted_answers.append({
                    "question": question,
                    "rating": None,
                    "answer_text": value,
                    "selected_option": None,
                })

            # ------------------------------------------------
            # UNKNOWN TYPE
            # ------------------------------------------------

            else:

                validation_errors.append(
                    f"Question {question.display_order}: "
                    "unsupported question type."
                )

        # ----------------------------------------------------
        # VALIDATION FAILED
        # ----------------------------------------------------

        if validation_errors:

            return render(
                request,
                "feedback/student_feedback_form.html",
                {
                    "assignment": assignment,
                    "questions": questions,
                    "errors": validation_errors,
                },
            )

        # ----------------------------------------------------
        # SAVE RESPONSE + ANSWERS
        # ----------------------------------------------------

        with transaction.atomic():

            response = (
                FeedbackResponse.objects.create(
                    assignment=assignment,
                    student=student,
                )
            )

            FeedbackAnswer.objects.bulk_create([

                FeedbackAnswer(
                    response=response,
                    question=item["question"],
                    rating=item["rating"],
                    answer_text=item[
                        "answer_text"
                    ],
                    selected_option=item[
                        "selected_option"
                    ],
                )

                for item in submitted_answers
            ])

        return redirect(
            "student_feedback_success",
            assignment_id=assignment.id,
        )

    # GET
    return render(
        request,
        "feedback/student_feedback_form.html",
        {
            "assignment": assignment,
            "questions": questions,
        },
    )


# ============================================================
# STUDENT FEEDBACK SUCCESS
# ============================================================

@student_required
def student_feedback_success(
    request,
    assignment_id,
):

    student = get_object_or_404(
        StudentProfile,
        user=request.user,
    )

    assignment = get_object_or_404(
        FeedbackAssignment,
        pk=assignment_id,
    )

    response = get_object_or_404(
        FeedbackResponse,
        assignment=assignment,
        student=student,
    )

    return render(
        request,
        "feedback/student_feedback_success.html",
        {
            "assignment": assignment,
            "response": response,
        },
    )
# ============================================================
# ADMIN FEEDBACK REPORTS
# ============================================================

@admin_required
def feedback_reports(request):

    assignments = (
        FeedbackAssignment.objects
        .select_related(
            "feedback_form",
            "course",
            "subject",
            "faculty",
            "faculty__user",
        )
        .order_by(
            "-created_at"
        )
    )

    selected_assignment_id = request.GET.get(
        "assignment"
    )

    selected_assignment = None

    if selected_assignment_id:
        selected_assignment = (
            FeedbackAssignment.objects
            .select_related(
                "feedback_form",
                "course",
                "subject",
                "faculty",
                "faculty__user",
            )
            .filter(
                pk=selected_assignment_id
            )
            .first()
        )

    responses_queryset = FeedbackResponse.objects.select_related(
        "assignment",
        "assignment__feedback_form",
        "assignment__course",
        "assignment__subject",
        "assignment__faculty",
        "assignment__faculty__user",
        "student",
        "student__user",
    )

    if selected_assignment:
        responses_queryset = responses_queryset.filter(
            assignment=selected_assignment
        )

    total_responses = responses_queryset.count()

    total_students = (
        responses_queryset
        .values("student_id")
        .distinct()
        .count()
    )

    # --------------------------------------------------------
    # QUESTION-WISE REPORT
    # --------------------------------------------------------

    question_reports = []

    questions = (
        FeedbackQuestion.objects
        .filter(
            feedback_form__assignments__in=(
                [selected_assignment]
                if selected_assignment
                else assignments
            ),
            is_active=True,
        )
        .distinct()
        .order_by(
            "feedback_form_id",
            "display_order",
            "id",
        )
    )

    for question in questions:

        answers = FeedbackAnswer.objects.filter(
            question=question,
            response__in=responses_queryset,
        )

        answer_count = answers.count()

        average_rating = None

        if question.question_type == "rating":

            rating_values = list(
                answers
                .exclude(rating__isnull=True)
                .values_list(
                    "rating",
                    flat=True,
                )
            )

            if rating_values:

                average_rating = round(
                    sum(rating_values)
                    / len(rating_values),
                    2,
                )

        option_counts = {}

        if question.question_type in {
            "choice",
            "yes_no",
        }:

            options = [
                question.option1,
                question.option2,
                question.option3,
                question.option4,
            ]

            options = [
                option
                for option in options
                if option
            ]

            if question.question_type == "yes_no":
                options = [
                    "yes",
                    "no",
                ]

            for option in options:

                option_counts[option] = (
                    answers
                    .filter(
                        selected_option=option
                    )
                    .count()
                )

        question_reports.append({
            "question": question,
            "answer_count": answer_count,
            "average_rating": average_rating,
            "option_counts": option_counts,
        })

    # --------------------------------------------------------
    # FACULTY-WISE REPORT
    # --------------------------------------------------------

    faculty_reports = {}

    for response in responses_queryset:

        faculty = response.assignment.faculty

        if faculty is None:
            continue

        faculty_id = faculty.id

        if faculty_id not in faculty_reports:

            faculty_name = (
                faculty.user.get_full_name().strip()
                if faculty.user.get_full_name()
                else faculty.user.username
            )

            faculty_reports[faculty_id] = {
                "faculty": faculty,
                "faculty_name": faculty_name,
                "responses": 0,
                "rating_values": [],
            }

        faculty_reports[faculty_id][
            "responses"
        ] += 1

        rating_values = (
            FeedbackAnswer.objects
            .filter(
                response=response,
                rating__isnull=False,
            )
            .values_list(
                "rating",
                flat=True,
            )
        )

        faculty_reports[
            faculty_id
        ]["rating_values"].extend(
            list(rating_values)
        )

    faculty_report_list = []

    for item in faculty_reports.values():

        rating_values = item.pop(
            "rating_values"
        )

        if rating_values:

            item["average_rating"] = round(
                sum(rating_values)
                / len(rating_values),
                2,
            )

        else:

            item["average_rating"] = None

        faculty_report_list.append(
            item
        )

    faculty_report_list.sort(
        key=lambda item: (
            item["average_rating"]
            if item["average_rating"] is not None
            else 0
        ),
        reverse=True,
    )

    # --------------------------------------------------------
    # OVERALL AVERAGE
    # --------------------------------------------------------

    all_ratings = list(
        FeedbackAnswer.objects
        .filter(
            response__in=responses_queryset,
            rating__isnull=False,
        )
        .values_list(
            "rating",
            flat=True,
        )
    )

    overall_average = None

    if all_ratings:

        overall_average = round(
            sum(all_ratings)
            / len(all_ratings),
            2,
        )

    # --------------------------------------------------------
    # RECENT RESPONSES
    # --------------------------------------------------------

    recent_responses = (
        responses_queryset
        .order_by(
            "-submitted_at"
        )[:50]
    )

    context = {
        "assignments": assignments,
        "selected_assignment": selected_assignment,

        "total_responses": total_responses,
        "total_students": total_students,
        "overall_average": overall_average,

        "question_reports": question_reports,
        "faculty_reports": faculty_report_list,
        "recent_responses": recent_responses,
    }

    return render(
        request,
        "feedback/reports.html",
        context,
    )


@admin_required
def export_feedback_reports(request):

    assignment_id = request.GET.get(
        "assignment"
    )

    responses = (
        FeedbackResponse.objects
        .select_related(
            "assignment",
            "assignment__feedback_form",
            "assignment__course",
            "assignment__subject",
            "assignment__faculty",
            "assignment__faculty__user",
            "student",
            "student__user",
        )
        .order_by(
            "submitted_at"
        )
    )

    if assignment_id:

        responses = responses.filter(
            assignment_id=assignment_id
        )

    response = HttpResponse(
        content_type="text/csv"
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; '
        'filename="feedback_report.csv"'
    )

    import csv

    writer = csv.writer(response)

    writer.writerow([
        "Student Username",
        "Student Name",
        "Feedback Form",
        "Course",
        "Semester",
        "Subject",
        "Faculty",
        "Question",
        "Question Type",
        "Rating",
        "Selected Option",
        "Text Answer",
        "Submitted At",
    ])

    for feedback_response in responses:

        student_name = (
            feedback_response.student.user
            .get_full_name()
            .strip()
        )

        if not student_name:
            student_name = (
                feedback_response.student.user.username
            )

        faculty_name = ""

        if feedback_response.assignment.faculty:

            faculty_name = (
                feedback_response
                .assignment
                .faculty
                .user
                .get_full_name()
                .strip()
            )

            if not faculty_name:

                faculty_name = (
                    feedback_response
                    .assignment
                    .faculty
                    .user
                    .username
                )

        answers = (
            FeedbackAnswer.objects
            .select_related(
                "question"
            )
            .filter(
                response=feedback_response
            )
            .order_by(
                "question__display_order",
                "question__id",
            )
        )

        for answer in answers:

            writer.writerow([
                feedback_response.student.user.username,
                student_name,
                feedback_response.assignment.feedback_form.title,
                feedback_response.assignment.course.name,
                feedback_response.assignment.semester,
                (
                    feedback_response.assignment.subject.name
                    if feedback_response.assignment.subject
                    else ""
                ),
                faculty_name,
                answer.question.question_text,
                answer.question.question_type,
                answer.rating
                if answer.rating is not None
                else "",
                answer.selected_option or "",
                answer.answer_text or "",
                feedback_response.submitted_at.strftime(
                    "%Y-%m-%d %H:%M:%S"
                ),
            ])

    return response
@admin_required
def bulk_create_feedback_assignments(request):

    from quiz.models import Course, FacultyProfile, Subject

    feedback_forms = FeedbackForm.objects.filter(
        is_active=True
    ).order_by("title")

    courses = Course.objects.all().order_by("name")

    if request.method == "POST":

        feedback_form_id = request.POST.get(
            "feedback_form"
        )

        course_id = request.POST.get(
            "course"
        )

        semester = request.POST.get(
            "semester"
        )

        academic_year = request.POST.get(
            "academic_year"
        )

        if not feedback_form_id:
            messages.error(
                request,
                "Please select a Feedback Form."
            )

            return redirect(
                "bulk_create_feedback_assignments"
            )

        if not course_id:
            messages.error(
                request,
                "Please select a Course."
            )

            return redirect(
                "bulk_create_feedback_assignments"
            )

        if not semester:
            messages.error(
                request,
                "Please select a Semester."
            )

            return redirect(
                "bulk_create_feedback_assignments"
            )

        if not academic_year:
            messages.error(
                request,
                "Please enter Academic Year."
            )

            return redirect(
                "bulk_create_feedback_assignments"
            )

        feedback_form = get_object_or_404(
            FeedbackForm,
            pk=feedback_form_id,
            is_active=True,
        )

        course = get_object_or_404(
            Course,
            pk=course_id,
        )

        semester = int(semester)

        subjects = Subject.objects.filter(
            course=course,
            semester=semester,
        ).distinct()

        faculties = (
            FacultyProfile.objects
            .filter(
                course=course,
                subjects__semester=semester,
            )
            .distinct()
            .select_related("user")
        )

        if not faculties.exists():

            messages.error(
                request,
                "No teaching faculty were found for the selected course and semester."
            )

            return redirect(
                "bulk_create_feedback_assignments"
            )

        created_count = 0
        skipped_count = 0

        with transaction.atomic():

            for faculty in faculties:

                faculty_subjects = subjects.filter(
                    facultyprofile__id=faculty.id
                )

                for subject in faculty_subjects:

                    exists = FeedbackAssignment.objects.filter(
                        feedback_form=feedback_form,
                        course=course,
                        semester=semester,
                        subject=subject,
                        faculty=faculty,
                        academic_year=academic_year,
                    ).exists()

                    if exists:

                        skipped_count += 1
                        continue

                    FeedbackAssignment.objects.create(
                        feedback_form=feedback_form,
                        course=course,
                        semester=semester,
                        subject=subject,
                        faculty=faculty,
                        academic_year=academic_year,
                        is_active=True,
                    )

                    created_count += 1

        messages.success(
            request,
            f"{created_count} faculty feedback assignments created successfully."
        )

        if skipped_count:

            messages.info(
                request,
                f"{skipped_count} existing assignments were skipped."
            )

        return redirect(
            "feedback_assignments"
        )

    return render(
        request,
        "feedback/bulk_create_assignment.html",
        {
            "feedback_forms": feedback_forms,
            "courses": courses,
        },
    )